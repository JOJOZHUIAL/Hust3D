# -*- coding: utf-8 -*-
"""申请表 Word 导出：按「附件1 3D打印服务申请表单」官方版式生成 docx。

版式要点（与学校下发样表保持一致）：
  A4 纵向，页边距上下 2.54cm、左右 3.17cm；
  标题「哈尔滨理工大学 / 3D打印服务申请表单」宋体 26pt 居中；
  四个部分标题加粗 12pt，正文宋体 12pt，填空处为下划线；
  勾选项用 □/☑，末尾落款宋体 7.5pt。
"""
import io
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

# 工作室收件邮箱（样表第三部分提示行与落款使用）
STUDIO_EMAIL = "hrbust3d101@163.com"

# 第四部分服务声明条款（与官方样表文字一致）
DOC_TERMS = [
    "1.我确认所提交的模型文件不涉及任何商业用途及侵权行为。",
    "2.我理解并同意，打印成品可能存在正常的层纹等工艺痕迹。",
    "3.我承诺提供的模型文件是完整且可打印的。我知晓如因模型自身问题（如破面、非流形、结构不合理）导致打印失败或效果不佳，责任由本人承担。",
    "4.我同意工作室根据《打印任务验收标准》对我的打印成品进行质量评判。",
    "5.我授权工作室为提供服务之目的，在我知情下处理我的模型文件。",
]

# 打印用途选项（value 与数据库 purpose 字段一致）
DOC_PURPOSES = [
    ("course", "课程作业"),
    ("research", "科研项目"),
    ("competition", "学科竞赛"),
    ("graduation", "毕业设计"),
    ("club", "学生社团项目"),
]


def _run(p, text, size=12, bold=False, underline=False, font="宋体"):
    """向段落追加一个 run 并设置中英文字体（正文默认宋体 12pt）。"""
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.underline = underline
    r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    return r


def _para(doc, align=None, space_after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    return p


def _blank_line(doc):
    _para(doc, space_after=0)


def _fill(value, pad=10):
    """把字段值包在空格里铺成下划线填空样式；空值退化为纯空白下划线。

    空白下划线依赖 _apply_compat 注入的 ulTrailSpace 标记才会被 Word 渲染。
    """
    text = str(value).strip() if value else ""
    if not text:
        return " " * pad
    return f"  {text}  "


def _blank(n):
    """n 个空格组成的空白下划线段。"""
    return " " * n


def _checkbox(doc, checked, label):
    """一行勾选项：☑/□ + 标签。"""
    p = _para(doc)
    _run(p, "☑" if checked else "□")
    _run(p, label)
    return p


def _underline_line(doc, label, value, value_size=12):
    """一行「标签：____」填空，值落在下划线上。"""
    p = _para(doc)
    _run(p, label)
    _run(p, _fill(value), underline=True)
    return p


def _signature_path(app_obj, upload_dir):
    """签名图片绝对路径；sign_url 形如 uploads/sign/202609/xxx.png。"""
    if not app_obj.sign_url:
        return None
    rel = app_obj.sign_url.split("uploads/", 1)[-1]
    path = os.path.join(upload_dir, rel)
    return path if os.path.isfile(path) else None


def _apply_compat(doc):
    """注入官方样表同款 Word 兼容标记（settings.xml /w:compat）。

    样表的空白填空是「下划线 + 空格」，Word 默认不为行尾空格绘制下划线，
    依赖 ulTrailSpace 等标记才会渲染；缺省生成的文档必须补上，否则
    无签名 / 无备注时的下划线空位在 Word 中不可见。
    旧式 compat 布尔元素需排在 compatSetting 之前，故依次插在 compat 首位。
    """
    settings = doc.settings.element
    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = settings.makeelement(qn("w:compat"), {})
        settings.append(compat)
    for name in ("spaceForUL", "balanceSingleByteDoubleByteWidth",
                 "doNotLeaveBackslashAlone", "ulTrailSpace",
                 "doNotExpandShiftReturn", "adjustLineHeightInTable", "useFELayout"):
        if compat.find(qn(f"w:{name}")) is None:
            compat.insert(0, compat.makeelement(qn(f"w:{name}"), {}))


def generate_application_docx(app_obj, upload_dir):
    """根据申请记录生成官方格式申请表，返回 BytesIO（docx 字节流）。"""
    user = app_obj.user
    name = user.name if user else ""
    college = user.college if user else ""
    phone = (user.phone if user else "") or ""
    email = (user.email if user else "") or ""

    created = app_obj.created_at
    date_str = f"{created.year} 年 {created.month} 月 {created.day} 日"

    doc = Document()

    # 与样表一致的兼容设置（行尾空格绘制下划线）
    _apply_compat(doc)

    # A4 页面与样表页边距
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(3.17)

    # 标题区
    _run(_para(doc, space_after=0), "附件1", 16, font="黑体")
    _blank_line(doc)
    _run(_para(doc, WD_ALIGN_PARAGRAPH.CENTER, 0), "哈尔滨理工大学", 26)
    _run(_para(doc, WD_ALIGN_PARAGRAPH.CENTER, 12), "3D打印服务申请表单", 26)
    p = _para(doc)
    _run(p, "申请日期：")
    _run(p, f"{created.year} 年 {created.month} 月 {created.day} 日")

    # 第一部分：申请人基本信息
    _run(_para(doc), "第一部分：申请人基本信息", bold=True)
    _underline_line(doc, "姓    名：", name)
    _underline_line(doc, "所属学院：", college)
    _underline_line(doc, "联系方式：", phone)
    _underline_line(doc, "电子邮箱：", email)
    _blank_line(doc)

    # 第二部分：打印需求详情
    _run(_para(doc), "第二部分：打印需求详情", bold=True)
    p = _para(doc, space_after=0)
    _run(p, "1.打印用途")
    for value, label in DOC_PURPOSES:
        _checkbox(doc, app_obj.purpose == value, label)
    # 其他用途：勾选时说明文字落在下划线上
    p = _para(doc)
    _run(p, "☑" if app_obj.purpose == "other" else "□")
    _run(p, "其他（请简要说明）：")
    _run(p, _fill(app_obj.purpose_other), underline=True)

    p = _para(doc, space_after=0)
    _run(p, "2.期望材料情况（工作室将根据库存情况进行调整）")
    _checkbox(doc, (app_obj.material or "").upper().startswith("PLA"), "PLA（FDM）")
    _checkbox(doc, False, "其他材料暂未开放")

    p = _para(doc)
    _run(p, "3.模型数量：")
    _run(p, _fill(app_obj.model_count, pad=8), underline=True)
    _run(p, " 个")
    _blank_line(doc)

    # 第三部分：模型文件与特别说明
    _run(_para(doc), "第三部分：模型文件与特别说明", bold=True)
    p = _para(doc)
    _run(p, "请将STL格式的模型文件，与此表单一并放送至邮箱：", bold=True, underline=True)
    _run(p, STUDIO_EMAIL, bold=True, underline=True)
    p = _para(doc, space_after=0)
    _run(p, "特别说明/备注：（如是否需要特定颜色等）")
    remark = (app_obj.remark or "").strip()
    if remark:
        for line in remark.splitlines():
            p = _para(doc)
            _run(p, _fill(line, pad=2) if line.strip() else _blank(60), underline=True)
    else:
        for _ in range(2):
            p = _para(doc)
            _run(p, _blank(60), underline=True)
    _blank_line(doc)

    # 第四部分：服务声明与确认
    _run(_para(doc), "第四部分：服务声明与确认", bold=True)
    p = _para(doc)
    _run(p, "我已阅读并同意以下条款：")
    for term in DOC_TERMS:
        _run(_para(doc), term)

    # 申请人签字：有电子签名则嵌入签名图片，否则保留下划线空位
    p = _para(doc)
    _run(p, "申请人签字：")
    sign_path = _signature_path(app_obj, upload_dir)
    if sign_path:
        try:
            p.add_run().add_picture(sign_path, height=Cm(1.0))
        except Exception:
            _run(p, _fill(None), underline=True)
    else:
        _run(p, _fill(None), underline=True)

    p = _para(doc)
    _run(p, "提交日期：")
    _run(p, f"{created.year} 年 {created.month} 月 {created.day} 日")

    # 落款（7.5pt 小字，与样表一致）
    _blank_line(doc)
    for text in (
        "表单提交后，将在24小时通过邮箱与您联络，请您注意查收。",
        "感谢您对工作室的支持！",
        f"“理工智造共享平台”管理团队{STUDIO_EMAIL}",
    ):
        _run(_para(doc, space_after=0), text, 7.5)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
